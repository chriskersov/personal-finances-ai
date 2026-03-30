import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx
from datetime import datetime
import ollama
import os

# ── Config ───────────────────────────────────────────────────────────────────
st.set_page_config(page_title="Finance AI", page_icon="₿", layout="centered")

MODEL = "qwen2.5:7b"
EXCEL_PATH = os.environ.get("FINANCE_FILE", "finances.xlsm")

# ── Styling ───────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Mono:wght@300;400;500&family=DM+Sans:wght@300;400;500&display=swap');
html, body, [data-testid="stAppViewContainer"] {
    background: #0a0a0f !important;
    color: #e8e4dc !important;
    font-family: 'DM Sans', sans-serif !important;
}
[data-testid="stHeader"], #MainMenu, footer { display: none !important; }
.main .block-container { padding: 3rem 2rem !important; max-width: 720px !important; }
.app-title {
    font-family: 'DM Mono', monospace;
    font-size: 0.75rem;
    letter-spacing: 0.25em;
    color: #5fffb0;
    text-transform: uppercase;
    margin-bottom: 0.25rem;
}
.app-month {
    font-family: 'DM Mono', monospace;
    font-size: 2rem;
    font-weight: 300;
    color: #e8e4dc;
    margin-bottom: 2rem;
}
.digest-box {
    background: #0f0f1a;
    border: 1px solid #1e1e35;
    border-left: 3px solid #5fffb0;
    border-radius: 4px;
    padding: 1.5rem 1.75rem;
    font-family: 'DM Mono', monospace;
    font-size: 0.82rem;
    line-height: 2;
    color: #a8a4b8;
    white-space: pre-wrap;
}
div[data-testid="stButton"] button {
    background: transparent !important;
    color: #444466 !important;
    border: 1px solid #1e1e35 !important;
    font-family: 'DM Mono', monospace !important;
    font-size: 0.7rem !important;
    letter-spacing: 0.1em !important;
    border-radius: 3px !important;
    margin-top: 1rem !important;
}
div[data-testid="stButton"] button:hover {
    color: #5fffb0 !important;
    border-color: #5fffb0 !important;
}
</style>
""", unsafe_allow_html=True)


# ── Data Extraction ───────────────────────────────────────────────────────────
def c(ws, col_letter, row):
    v = ws.cell(row, col_idx(col_letter)).value
    if v is None or v == "" or str(v) == "#DIV/0!":
        return None
    return v

def fmt(v):
    if v is None: return "—"
    try: return f"£{float(v):,.2f}"
    except: return str(v)

def fmt_pct(v):
    if v is None: return "—"
    try: return f"{float(v)*100:.0f}%"
    except: return str(v)


@st.cache_data(ttl=30)
def load_sheet(path, sheet_name):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    # ── Needs (B21:E100, data starts row 24)
    needs = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'B',r), c(ws,'C',r), c(ws,'D',r), c(ws,'E',r)
        if actual is not None:
            needs.append({"description": desc, "date": str(date)[:10] if date else None,
                          "category": cat, "actual": actual})

    # ── Wants (G21:J100, data starts row 24)
    wants = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'G',r), c(ws,'H',r), c(ws,'I',r), c(ws,'J',r)
        if actual is not None:
            wants.append({"description": desc, "date": str(date)[:10] if date else None,
                          "category": cat, "actual": actual})

    # ── Savings (L21:O100, data starts row 24)
    savings = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'L',r), c(ws,'M',r), c(ws,'N',r), c(ws,'O',r)
        if actual is not None:
            savings.append({"description": desc, "date": str(date)[:10] if date else None,
                            "category": cat, "actual": actual})

    # ── Cash Flow (Q43:W46 — rows 43=Needs, 44=Wants, 45=Savings, 46=Total)
    cash_flow = {}
    for row, label in [(43,"needs"),(44,"wants"),(45,"savings"),(46,"total")]:
        cash_flow[label] = {
            "in":   c(ws,'U',row),
            "out":  c(ws,'V',row),
            "diff": c(ws,'W',row),
        }

    # ── Budget goals vs actual (Q51:V53 — rows 51=Needs, 52=Wants, 53=Savings)
    goals = {}
    for row, label in [(51,"needs"),(52,"wants"),(53,"savings")]:
        goals[label] = {
            "goal":   c(ws,'T',row),
            "actual": c(ws,'V',row),
        }

    # ── Income (Q58:V100, desc=Q, date=T, actual=V)
    income = []
    for r in range(58, 101):
        desc, date, actual = c(ws,'Q',r), c(ws,'T',r), c(ws,'V',r)
        if actual is not None:
            income.append({"description": desc,
                           "date": str(date)[:10] if date else None,
                           "actual": actual})

    # ── Categories (AA53:AE66)
    categories = []
    for r in range(53, 67):
        name = c(ws,'AA',r)
        if name:
            categories.append({
                "category": name,
                "needs":    c(ws,'AB',r),
                "wants":    c(ws,'AC',r),
                "savings":  c(ws,'AD',r),
                "total":    c(ws,'AE',r),
            })

    return {
        "needs": needs,
        "wants": wants,
        "savings": savings,
        "cash_flow": cash_flow,
        "goals": goals,
        "income": income,
        "categories": categories,
    }


def build_context(sheet_name, d):
    lines = []
    lines.append(f"MONTH: {sheet_name}")
    lines.append(f"Data read on: {datetime.now().strftime('%d %B %Y')}\n")

    # Cash flow
    cf = d["cash_flow"]
    lines.append("CASH FLOW:")
    for label in ["needs","wants","savings","total"]:
        v = cf[label]
        lines.append(f"  {label.capitalize():8s} — In: {fmt(v['in'])}  Out: {fmt(v['out'])}  Difference: {fmt(v['diff'])}")

    # Goals
    lines.append("\nBUDGET GOALS vs ACTUAL:")
    for label, v in d["goals"].items():
        lines.append(f"  {label.capitalize():8s} — Goal: {fmt_pct(v['goal'])}  Actual: {fmt_pct(v['actual'])}")

    # Income
    if d["income"]:
        lines.append(f"\nINCOME ({len(d['income'])} entries):")
        for e in d["income"]:
            lines.append(f"  {e['date']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nINCOME: none recorded yet")

    # Needs
    if d["needs"]:
        total = sum(float(e["actual"]) for e in d["needs"])
        lines.append(f"\nNEEDS SPENDING — {len(d['needs'])} entries, total {fmt(total)}:")
        for e in d["needs"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nNEEDS SPENDING: none recorded yet")

    # Wants
    if d["wants"]:
        total = sum(float(e["actual"]) for e in d["wants"])
        lines.append(f"\nWANTS SPENDING — {len(d['wants'])} entries, total {fmt(total)}:")
        for e in d["wants"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nWANTS SPENDING: none recorded yet")

    # Savings
    if d["savings"]:
        total = sum(float(e["actual"]) for e in d["savings"])
        lines.append(f"\nSAVINGS — {len(d['savings'])} entries, total {fmt(total)}:")
        for e in d["savings"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nSAVINGS: none recorded yet")

    # Categories
    active_cats = [cat for cat in d["categories"] if cat["total"] not in (None, 0)]
    if active_cats:
        lines.append(f"\nSPEND BY CATEGORY:")
        for cat in sorted(active_cats, key=lambda x: float(x["total"] or 0), reverse=True):
            lines.append(f"  {cat['category']:15s} — Needs: {fmt(cat['needs'])}  Wants: {fmt(cat['wants'])}  Total: {fmt(cat['total'])}")
    else:
        lines.append("\nSPEND BY CATEGORY: no data yet")

    return "\n".join(lines)


def generate_digest(context, sheet_name):
    now = datetime.now()
    day_of_month = now.day
    days_in_month = 31 if now.month in [1,3,5,7,8,10,12] else 30 if now.month in [4,6,9,11] else 28
    pct_through = round((day_of_month / days_in_month) * 100)

    prompt = f"""Write a monthly financial summary for {sheet_name}.
Today is {now.strftime('%d %B %Y')} — day {day_of_month} of {days_in_month} ({pct_through}% through the month).

Use only the numbers in the data. Be specific. Write 4-5 sentences covering:
- Total income vs total spending and the net position
- How the needs / wants / savings split compares to the goals (25% / 15% / 60%)
- The biggest spending category
- One honest, useful observation for the rest of the month

If any section shows no data or zeros, say so clearly.
Write as a plain paragraph. No bullet points. Calm, direct tone like a smart financial advisor."""

    response = ollama.chat(
        model=MODEL,
        messages=[
            {"role": "system", "content": f"You are a personal finance assistant. Here is the user's full finance data for this month:\n\n{context}"},
            {"role": "user", "content": prompt}
        ]
    )
    return response["message"]["content"]


# ── App ───────────────────────────────────────────────────────────────────────
now = datetime.now()

st.markdown('<div class="app-title">Finance AI</div>', unsafe_allow_html=True)
st.markdown(f'<div class="app-month">{now.strftime("%B %Y")}</div>', unsafe_allow_html=True)

if not os.path.exists(EXCEL_PATH):
    st.error(f"Excel file not found.\n\nSet the path by running:\nexport FINANCE_FILE=/full/path/to/your/file.xlsm")
    st.stop()

# Find the current month's sheet
try:
    from openpyxl import load_workbook as _lw
    _wb = _lw(EXCEL_PATH, read_only=True)
    sheet_names = _wb.sheetnames
    _wb.close()
except Exception as e:
    st.error(f"Could not open Excel file: {e}")
    st.stop()

# Match current month to sheet name
current_month_str = now.strftime("%B %Y")  # e.g. "March 2026"
sheet_name = next((s for s in sheet_names if current_month_str.lower() in s.lower()), sheet_names[-1])

try:
    data = load_sheet(EXCEL_PATH, sheet_name)
    context = build_context(sheet_name, data)
except Exception as e:
    st.error(f"Could not read sheet '{sheet_name}': {e}")
    st.stop()

if "digest" not in st.session_state:
    st.session_state.digest = None

if st.session_state.digest is None:
    with st.spinner("Reading your finances…"):
        st.session_state.digest = generate_digest(context, sheet_name)

st.markdown(f'<div class="digest-box">{st.session_state.digest}</div>', unsafe_allow_html=True)

if st.button("↻ Regenerate"):
    with st.spinner("Thinking…"):
        st.session_state.digest = generate_digest(context, sheet_name)
    st.rerun()