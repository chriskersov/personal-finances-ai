import os
import ollama
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx

from textual.app import App, ComposeResult
from textual.containers import ScrollableContainer
from textual.widgets import Header, Footer, Static, Button
from textual.worker import Worker, WorkerState

# ── Config ───────────────────────────────────────────────────────────────────
MODEL = "qwen2.5:7b"
EXCEL_PATH = os.environ.get("FINANCE_FILE", "finances.xlsm")

# ── Pure Logic ────────────────────────────────────────────────────────────────
def c(ws, col_letter, row):
    v = ws.cell(row, col_idx(col_letter)).value
    if v is None or v == "" or str(v) == "#DIV/0!":
        return None
    return v

def fmt(v):
    if v is None: return "—"
    try: return f"£{float(v):,.2f}"
    except: return str(v)

def fmt_pct(v):
    if v is None: return "—"
    try: return f"{float(v)*100:.0f}%"
    except: return str(v)


def load_sheet(path, sheet_name):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    # ── Needs (B21:E100, data starts row 24)
    needs = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'B',r), c(ws,'C',r), c(ws,'D',r), c(ws,'E',r)
        if actual is not None:
            needs.append({"description": desc, "date": str(date)[:10] if date else None,
                          "category": cat, "actual": actual})

    # ── Wants (G21:J100, data starts row 24)
    wants = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'G',r), c(ws,'H',r), c(ws,'I',r), c(ws,'J',r)
        if actual is not None:
            wants.append({"description": desc, "date": str(date)[:10] if date else None,
                          "category": cat, "actual": actual})

    # ── Savings (L21:O100, data starts row 24)
    savings = []
    for r in range(24, 101):
        desc, date, cat, actual = c(ws,'L',r), c(ws,'M',r), c(ws,'N',r), c(ws,'O',r)
        if actual is not None:
            savings.append({"description": desc, "date": str(date)[:10] if date else None,
                            "category": cat, "actual": actual})

    # ── Cash Flow (Q43:W46)
    cash_flow = {}
    for row, label in [(43,"needs"),(44,"wants"),(45,"savings"),(46,"total")]:
        cash_flow[label] = {
            "in":   c(ws,'U',row),
            "out":  c(ws,'V',row),
            "diff": c(ws,'W',row),
        }

    # ── Budget goals vs actual (Q51:V53)
    goals = {}
    for row, label in [(51,"needs"),(52,"wants"),(53,"savings")]:
        goals[label] = {
            "goal":   c(ws,'T',row),
            "actual": c(ws,'V',row),
        }

    # ── Income (Q58:V100)
    income = []
    for r in range(58, 101):
        desc, date, actual = c(ws,'Q',r), c(ws,'T',r), c(ws,'V',r)
        if actual is not None:
            income.append({"description": desc,
                           "date": str(date)[:10] if date else None,
                           "actual": actual})

    # ── Categories (AA53:AE66)
    categories = []
    for r in range(53, 67):
        name = c(ws,'AA',r)
        if name:
            categories.append({
                "category": name,
                "needs":    c(ws,'AB',r),
                "wants":    c(ws,'AC',r),
                "savings":  c(ws,'AD',r),
                "total":    c(ws,'AE',r),
            })

    return {
        "needs": needs,
        "wants": wants,
        "savings": savings,
        "cash_flow": cash_flow,
        "goals": goals,
        "income": income,
        "categories": categories,
    }


def build_context(sheet_name, d):
    lines = []
    lines.append(f"MONTH: {sheet_name}")
    lines.append(f"Data read on: {datetime.now().strftime('%d %B %Y')}\n")

    # Cash flow
    cf = d["cash_flow"]
    lines.append("CASH FLOW:")
    for label in ["needs","wants","savings","total"]:
        v = cf[label]
        lines.append(f"  {label.capitalize():8s} — In: {fmt(v['in'])}  Out: {fmt(v['out'])}  Difference: {fmt(v['diff'])}")

    # Goals
    lines.append("\nBUDGET GOALS vs ACTUAL:")
    for label, v in d["goals"].items():
        lines.append(f"  {label.capitalize():8s} — Goal: {fmt_pct(v['goal'])}  Actual: {fmt_pct(v['actual'])}")

    # Income
    if d["income"]:
        lines.append(f"\nINCOME ({len(d['income'])} entries):")
        for e in d["income"]:
            lines.append(f"  {e['date']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nINCOME: none recorded yet")

    # Needs
    if d["needs"]:
        total = sum(float(e["actual"]) for e in d["needs"])
        lines.append(f"\nNEEDS SPENDING — {len(d['needs'])} entries, total {fmt(total)}:")
        for e in d["needs"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nNEEDS SPENDING: none recorded yet")

    # Wants
    if d["wants"]:
        total = sum(float(e["actual"]) for e in d["wants"])
        lines.append(f"\nWANTS SPENDING — {len(d['wants'])} entries, total {fmt(total)}:")
        for e in d["wants"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nWANTS SPENDING: none recorded yet")

    # Savings
    if d["savings"]:
        total = sum(float(e["actual"]) for e in d["savings"])
        lines.append(f"\nSAVINGS — {len(d['savings'])} entries, total {fmt(total)}:")
        for e in d["savings"]:
            lines.append(f"  {e['date']} | {e['category']} | {e['description']} | {fmt(e['actual'])}")
    else:
        lines.append("\nSAVINGS: none recorded yet")

    # Categories
    active_cats = [cat for cat in d["categories"] if cat["total"] not in (None, 0)]
    if active_cats:
        lines.append(f"\nSPEND BY CATEGORY:")
        for cat in sorted(active_cats, key=lambda x: float(x["total"] or 0), reverse=True):
            lines.append(f"  {cat['category']:15s} — Needs: {fmt(cat['needs'])}  Wants: {fmt(cat['wants'])}  Total: {fmt(cat['total'])}")
    else:
        lines.append("\nSPEND BY CATEGORY: no data yet")

    return "\n".join(lines)


def generate_digest(context, sheet_name):
    now = datetime.now()
    day_of_month = now.day
    days_in_month = 31 if now.month in [1,3,5,7,8,10,12] else 30 if now.month in [4,6,9,11] else 28
    pct_through = round((day_of_month / days_in_month) * 100)

    prompt = f"""Write a monthly financial summary for {sheet_name}.
Today is {now.strftime('%d %B %Y')} — day {day_of_month} of {days_in_month} ({pct_through}% through the month).

Use only the numbers in the data. Be specific. Write 4-5 sentences covering:
- Total income vs total spending and the net position
- How the needs / wants / savings split compares to the goals (25% / 15% / 60%)
- The biggest spending category
- One honest, useful observation for the rest of the month

If any section shows no data or zeros, say so clearly.
Write as a plain paragraph. No bullet points. Calm, direct tone like a smart financial advisor."""

    response = ollama.chat(
        model=MODEL,
        messages=[
            {"role": "system", "content": f"You are a personal finance assistant. Here is the user's full finance data for this month:\n\n{context}"},
            {"role": "user", "content": prompt}
        ]
    )
    return response["message"]["content"]


# ── TUI Application ──────────────────────────────────────────────────────────
class FinanceApp(App):
    CSS = """
    Screen { background: #0e0e0e; color: #e8e4dc; }

    #title {
        color: #5fffb0;
        text-style: bold;
        margin: 1 0 0 2;
    }

    #month {
        margin: 0 0 0 2;
        text-style: bold;
        color: #e8e4dc;
    }

    #divider {
        color: #1e3d2e;
        margin: 0 2;
    }

    #digest-box {
        background: #0e0e0e;
        border: solid #1e3d2e;
        border-left: tall #5fffb0;
        padding: 1 2;
        margin: 1 2 1 2;
        height: auto;
        color: #c8c4bc;
    }

    #digest-box.loading {
        color: #5fffb0;
        text-style: italic;
    }

    #regen {
        background: #0e1a14;
        color: #5fffb0;
        border: solid #1e3d2e;
        margin-left: 2;
        margin-bottom: 1;
    }

    #regen:hover {
        background: #1a3d28;
        color: #5fffb0;
        border: solid #5fffb0;
    }

    #regen:disabled {
        background: #0e0e0e;
        color: #2a4a38;
        border: solid #1a2a20;
    }

    Footer {
        background: #0e1a14;
        color: #2a6644;
    }

    #error-box {
        background: #180f0f;
        border: solid #3a1a1a;
        border-left: tall #ff5f5f;
        padding: 1 2;
        margin: 1 2 1 2;
        height: auto;
        color: #ff9a9a;
    }
    """

    def compose(self) -> ComposeResult:
        yield Static("FINANCE AI", id="title")
        yield Static(datetime.now().strftime("%B %Y"), id="month")
        yield Static("─" * 80, id="divider")
        with ScrollableContainer():
            yield Static("Reading your finances…", id="digest-box", classes="loading")
            yield Button("↻  REGENERATE", id="regen")
        yield Footer()

    def on_mount(self) -> None:
        if not os.path.exists(EXCEL_PATH):
            self._show_error(
                f"Excel file not found at: {EXCEL_PATH}\n\n"
                f"Set the path with:\n  export FINANCE_FILE=/full/path/to/your/file.xlsm"
            )
            return
        self._run_digest()

    def _run_digest(self) -> None:
        box = self.query_one("#digest-box", Static)
        box.update("Reading your finances…")
        box.add_class("loading")
        self.query_one("#regen", Button).disabled = True
        self.run_worker(self._worker_logic, thread=True, exclusive=True)

    def _worker_logic(self) -> str:
        now = datetime.now()

        wb = load_workbook(EXCEL_PATH, read_only=True)
        sheet_names = wb.sheetnames
        wb.close()

        current_month_str = now.strftime("%B %Y")
        sheet_name = next(
            (s for s in sheet_names if current_month_str.lower() in s.lower()),
            sheet_names[-1]
        )

        data = load_sheet(EXCEL_PATH, sheet_name)
        context = build_context(sheet_name, data)
        return generate_digest(context, sheet_name)

    def on_worker_state_changed(self, event: Worker.StateChanged) -> None:
        box = self.query_one("#digest-box", Static)
        btn = self.query_one("#regen", Button)

        if event.state == WorkerState.SUCCESS:
            box.remove_class("loading")
            box.update(event.worker.result)
            btn.disabled = False

        elif event.state == WorkerState.ERROR:
            box.remove_class("loading")
            self._show_error(f"Error: {event.worker.error}")
            btn.disabled = False

    def on_button_pressed(self, event: Button.Pressed) -> None:
        if event.button.id == "regen":
            self._run_digest()

    def _show_error(self, message: str) -> None:
        box = self.query_one("#digest-box", Static)
        box.remove_class("loading")
        box.add_class("error-box")
        box.update(message)


if __name__ == "__main__":
    FinanceApp().run()